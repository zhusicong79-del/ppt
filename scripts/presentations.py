"""Private tenant-local OOXML templates. Never commit processed client content."""
import sys,json,zipfile,copy,posixpath,tempfile,subprocess,shutil,os
from pathlib import Path
from lxml import etree as E
from pptx import Presentation
import fitz
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
REL='http://schemas.openxmlformats.org/package/2006/relationships'; CT='http://schemas.openxmlformats.org/package/2006/content-types'
def xml(b):return E.fromstring(b,E.XMLParser(resolve_entities=False,no_network=True))
def dump(e):return E.tostring(e,xml_declaration=True,encoding='UTF-8',standalone=True)
def pdf_for(pptx,out):
 out.mkdir(parents=True,exist_ok=True)
 with tempfile.TemporaryDirectory() as profile:
  r=subprocess.run([os.environ.get('SOFFICE','soffice'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf:impress_pdf_Export:{"ExportHiddenSlides":{"type":"boolean","value":"true"}}','--outdir',str(out),str(pptx)],capture_output=True,timeout=300)
  pdf=out/(pptx.stem+'.pdf')
  if r.returncode or not pdf.exists():raise RuntimeError('PPT conversion failed; check LibreOffice installation')
 return pdf
def text_shape(sh,w,h):
 if not sh.has_text_frame or not sh.text.strip():return None
 runs=[run for p in sh.text_frame.paragraphs for run in p.runs]
 sizes=[run.font.size.pt for run in runs if run.font.size]
 return {'id':str(sh.shape_id),'text':sh.text,'x':sh.left/w*100,'y':sh.top/h*100,'w':sh.width/w*100,'h':sh.height/h*100,'fontSize':max(sizes or [18]),'font':next((run.font.name for run in runs if run.font.name),'Microsoft Yahei')}
def import_file(root,key,name):
 folder=root/'templates';src=folder/(key+'.pptx')
 with zipfile.ZipFile(src) as z:
  if sum(i.file_size for i in z.infolist())>200000000 or any(i.file_size>80000000 for i in z.infolist()):raise ValueError('Expanded PPTX too large')
 r=Presentation(src);out=folder/key;out.mkdir(exist_ok=True)
 pdf=pdf_for(src,out);doc=fitz.open(pdf)
 if len(doc)!=len(r.slides):raise RuntimeError('Rendered page count differs from template')
 slides=[]
 for i,s in enumerate(r.slides):
  shapes=[t for sh in s.shapes if (t:=text_shape(sh,r.slide_width,r.slide_height))]
  meaningful=[t for t in shapes if len(t['text'].strip())>3 and t['text'].strip() not in ['SmartX','SMARTX']]
  title=max(meaningful,key=lambda t:t['fontSize'],default={'text':'页面 '+str(i+1),'id':None})
  bodies=[t for t in meaningful if t['id']!=title['id'] and len(t['text'])>12]
  body=max(bodies,key=lambda t:len(t['text']),default={'text':'','id':None})
  doc[i].get_pixmap(matrix=fitz.Matrix(1,1)).save(str(out/f'{i+1}.png'))
  slides.append({'id':f'{key}-page-{i+1}','title':title['text'],'body':body['text'],'tag':'公司原始模板','type':'features','category':'公司模板','template':key,'source':{'fileId':key,'fileName':name,'page':i+1},'native':{'templateId':key,'page':i+1,'titleShape':title['id'],'bodyShape':body['id'],'textShapes':shapes},'edits':{}})
 result={'id':key,'name':name,'pages':len(slides),'modified':'原始附件','category':'公司模板','folder':'公司正式 PPT','slides':slides,'official':True,'width':r.slide_width,'height':r.slide_height,'masters':len(r.slide_masters)}
 (out/'manifest.json').write_text(json.dumps(result,ensure_ascii=False))
 return {'id':key,'pages':len(slides),'masters':len(r.slide_masters)}
def set_text(sp,text):
 tx=sp.find('p:txBody',NS)
 if tx is None:return
 old=tx.findall('a:p',NS);ppr=old[0].find('a:pPr',NS) if old else None
 rpr=next(iter(tx.findall('.//a:rPr',NS)),None)
 for p in old:tx.remove(p)
 for line in text.split('\n'):
  p=E.SubElement(tx,'{'+NS['a']+'}p')
  if ppr is not None:p.append(copy.deepcopy(ppr))
  run=E.SubElement(p,'{'+NS['a']+'}r')
  if rpr is not None:run.append(copy.deepcopy(rpr))
  E.SubElement(run,'{'+NS['a']+'}t').text=line

def build(root,deck,out):
 slides=deck['slides'];ids=list(dict.fromkeys(s.get('native',{}).get('templateId') for s in slides))
 if None in ids or not slides:raise ValueError('Every page must be based on an uploaded company template')
 zips={k:zipfile.ZipFile(root/'templates'/(k+'.pptx')) for k in ids};first=zips[ids[0]]
 contents={};defaults={};overrides={}
 for key,z in zips.items():
  prefix='sources/'+key+'/'
  for name in z.namelist():
   if name=='[Content_Types].xml':continue
   contents[prefix+name]=z.read(name)
  types=xml(z.read('[Content_Types].xml'))
  for item in types:
   if E.QName(item).localname=='Default':defaults[item.get('Extension')]=item.get('ContentType')
   else:overrides['/'+prefix+item.get('PartName').lstrip('/')]=item.get('ContentType')
 pres=xml(first.read('ppt/presentation.xml')); rels=E.Element('{'+REL+'}Relationships',nsmap={None:REL})
 oldrels=xml(first.read('ppt/_rels/presentation.xml.rels'))
 for old in oldrels:
  if old.get('Type').endswith('/slide') or old.get('Type').endswith('/slideMaster'):continue
  new=copy.deepcopy(old)
  if old.get('TargetMode')!='External':new.set('Target','../sources/'+ids[0]+'/'+posixpath.normpath('ppt/'+old.get('Target')))
  rels.append(new)
 lst=pres.find('p:sldIdLst',NS)
 if lst is None:lst=E.SubElement(pres,'{'+NS['p']+'}sldIdLst')
 for c in list(lst):lst.remove(c)
 masters=pres.find('p:sldMasterIdLst',NS)
 if masters is None:masters=E.Element('{'+NS['p']+'}sldMasterIdLst');pres.insert(0,masters)
 for c in list(masters):masters.remove(c)
 mi=0
 for key,z in zips.items():
  pr=xml(z.read('ppt/_rels/presentation.xml.rels'))
  for rel in pr:
   if not rel.get('Type').endswith('/slideMaster'):continue
   mi+=1;rid='companyMaster'+str(mi)
   E.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=rel.get('Type'),Target='../sources/'+key+'/'+posixpath.normpath('ppt/'+rel.get('Target')))
   E.SubElement(masters,'{'+NS['p']+'}sldMasterId',id=str(2147483647+mi),attrib={'{'+NS['r']+'}id':rid})
 # Preserve embedded brand fonts from every source presentation.
 embedded=pres.find('p:embeddedFontLst',NS)
 if embedded is not None:
  for child in list(embedded):embedded.remove(child)
 else:
  embedded=E.SubElement(pres,'{'+NS['p']+'}embeddedFontLst')
 seenfonts=set();fontcounter=0
 for key,z in zips.items():
  sourcepres=xml(z.read('ppt/presentation.xml'));sourcerels={r.get('Id'):r for r in xml(z.read('ppt/_rels/presentation.xml.rels'))}
  for font in sourcepres.findall('p:embeddedFontLst/p:embeddedFont',NS):
   face=font.find('p:font',NS).get('typeface')
   if face in seenfonts:continue
   seenfonts.add(face);newfont=copy.deepcopy(font)
   for style in newfont:
    rid=style.get('{'+NS['r']+'}id')
    if not rid:continue
    old=sourcerels[rid];fontcounter+=1;newid='companyFont'+str(fontcounter);style.set('{'+NS['r']+'}id',newid)
    E.SubElement(rels,'{'+REL+'}Relationship',Id=newid,Type=old.get('Type'),Target='../sources/'+key+'/'+posixpath.normpath('ppt/'+old.get('Target')))
   embedded.append(newfont)
 for i,s in enumerate(slides,1):
  n=s['native'];key=n['templateId'];z=zips[key];page=int(n['page']);path=f'ppt/slides/slide{page}.xml';tree=xml(z.read(path));tree.attrib.pop('show',None)
  edits=s.get('edits',{}).copy()
  if n.get('titleShape') and s['title']!=next((t['text'] for t in n.get('textShapes',[]) if t['id']==str(n['titleShape'])),s['title']):edits[str(n['titleShape'])]=s['title']
  if n.get('bodyShape') and s.get('body') is not None and s['body']!=next((t['text'] for t in n.get('textShapes',[]) if t['id']==str(n['bodyShape'])),s['body']):edits[str(n['bodyShape'])]=s['body']
  for sp in tree.findall('.//p:sp',NS):
   prop=sp.find('p:nvSpPr/p:cNvPr',NS)
   if prop is not None and prop.get('id') in edits:set_text(sp,str(edits[prop.get('id')]))
  # Topology is exported as editable native PowerPoint shapes and connectors.
  if s.get('diagram'):add_diagram(tree,s['diagram'],pres.find('p:sldSz',NS))
  target=f'ppt/slides/slide{i}.xml';contents[target]=dump(tree)
  relpath=f'ppt/slides/_rels/slide{page}.xml.rels'
  if relpath in z.namelist():
   sr=xml(z.read(relpath))
   for rel in sr:
    if rel.get('TargetMode')!='External':rel.set('Target','/sources/'+key+'/'+posixpath.normpath(posixpath.join(posixpath.dirname(path),rel.get('Target'))))
   contents[f'ppt/slides/_rels/slide{i}.xml.rels']=dump(sr)
  E.SubElement(rels,'{'+REL+'}Relationship',Id='exportSlide'+str(i),Type=NS['r']+'/slide',Target='slides/slide'+str(i)+'.xml')
  E.SubElement(lst,'{'+NS['p']+'}sldId',id=str(255+i),attrib={'{'+NS['r']+'}id':'exportSlide'+str(i)})
  overrides['/'+target]='application/vnd.openxmlformats-officedocument.presentationml.slide+xml'
 contents['ppt/presentation.xml']=dump(pres);contents['ppt/_rels/presentation.xml.rels']=dump(rels)
 overrides['/ppt/presentation.xml']='application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml'
 rootrels=E.Element('{'+REL+'}Relationships',nsmap={None:REL});E.SubElement(rootrels,'{'+REL+'}Relationship',Id='office',Type=NS['r']+'/officeDocument',Target='ppt/presentation.xml');contents['_rels/.rels']=dump(rootrels)
 ct=E.Element('{'+CT+'}Types',nsmap={None:CT})
 for ext,kind in defaults.items():E.SubElement(ct,'{'+CT+'}Default',Extension=ext,ContentType=kind)
 for part,kind in overrides.items():E.SubElement(ct,'{'+CT+'}Override',PartName=part,ContentType=kind)
 contents['[Content_Types].xml']=dump(ct)
 # Keep only reachable package parts; do not embed unrelated internal slides.
 reachable={'_rels/.rels'};queue=['']
 while queue:
  part=queue.pop();rp=posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels') if part else '_rels/.rels'
  if rp not in contents:continue
  reachable.add(rp)
  for rel in xml(contents[rp]):
   if rel.get('TargetMode')=='External':continue
   target=rel.get('Target').split('#')[0]
   target=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part),target))
   if target in contents and target not in reachable:reachable.add(target);queue.append(target)
 for item in list(ct):
  if E.QName(item).localname=='Override' and item.get('PartName').lstrip('/') not in reachable:ct.remove(item)
 contents['[Content_Types].xml']=dump(ct);reachable.add('[Content_Types].xml')
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for path in reachable:z.writestr(path,contents[path])
 for z in zips.values():z.close()

def add_diagram(tree,d,size):
 # Use python-pptx to construct native, individually editable diagram objects.
 from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
 from pptx.dml.color import RGBColor
 from pptx.util import Pt
 r=Presentation();r.slide_width=int(size.get('cx'));r.slide_height=int(size.get('cy'));sl=r.slides.add_slide(r.slide_layouts[6]);w=r.slide_width;h=r.slide_height
 def xy(n):return int((.05+n['x']/1000*.9)*w),int((.29+n['y']/430*.63)*h)
 for e in d['edges']:
  a=next(n for n in d['nodes'] if n['id']==e['from']);b=next(n for n in d['nodes'] if n['id']==e['to']);ax,ay=xy(a);bx,by=xy(b)
  line=sl.shapes.add_connector(MSO_CONNECTOR.ELBOW if e['style']=='elbow' else MSO_CONNECTOR.STRAIGHT,ax+int(w*.04),ay+int(h*.055),bx+int(w*.04),by+int(h*.055));line.line.color.rgb=RGBColor.from_string('0080FF');line.line.width=Pt(2)
  if e['style']=='dashed':
   from pptx.enum.dml import MSO_LINE_DASH_STYLE
   line.line.dash_style=MSO_LINE_DASH_STYLE.DASH
  if e.get('arrow'):
   ln=line._element.find('.//a:ln',NS);E.SubElement(ln,'{'+NS['a']+'}tailEnd',type='triangle')
  if e.get('label'):
   box=sl.shapes.add_textbox((ax+bx)//2,(ay+by)//2-int(h*.03),int(w*.17),int(h*.04));box.text=e['label'];box.text_frame.paragraphs[0].font.size=Pt(10)
 for n in d['nodes']:
  x,y=xy(n);kind=MSO_SHAPE.CLOUD if n['type']=='cloud' else MSO_SHAPE.CAN if n['type']=='storage' else MSO_SHAPE.ROUNDED_RECTANGLE
  shape=sl.shapes.add_shape(kind,x,y,int(w*.09),int(h*.13));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('D8EAFF');shape.line.color.rgb=RGBColor.from_string('0080FF');shape.text=n['label'];shape.text_frame.paragraphs[0].font.size=Pt(12)
 sp=tree.find('p:cSld/p:spTree',NS);largest=max([int(x.get('id','0')) for x in tree.findall('.//p:cNvPr',NS)] or [0])
 for sh in sl.shapes:
  el=copy.deepcopy(sh._element)
  for prop in el.findall('.//p:cNvPr',NS):largest+=1;prop.set('id',str(largest))
  sp.append(el)

if __name__=='__main__':
 mode=sys.argv[1];root=Path(sys.argv[2])
 if mode=='import':print(json.dumps(import_file(root,sys.argv[3],sys.argv[4]),ensure_ascii=False))
 elif mode=='export':
  deck=json.loads(Path(sys.argv[3]).read_text());out=Path(sys.argv[4]);build(root,deck,out)
  if sys.argv[5]=='pptx':print(json.dumps({'file':str(out)}))
  else:
   pdf=pdf_for(out,out.parent);doc=fitz.open(pdf)
   if sys.argv[5]=='pdf':print(json.dumps({'file':str(pdf),'pages':len(doc)}))
   else:
    pages=[]
    for i,page in enumerate(doc):
     path=out.parent/f'page-{i+1}.png';page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(str(path));pages.append(str(path))
    print(json.dumps({'pages':pages,'pdf':str(pdf)}))
